"""
Generates the customer downloads from data/menu.json:
  public/downloads/daddus-biryani-menu.pdf
  public/downloads/daddus-biryani-menu.xlsx
Run after scripts/extract_menu.py whenever prices change.
Fonts: place RozhaOne-Regular.ttf and Mukta-*.ttf in scripts/fonts/
"""
import json, os
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph, Table,
                                TableStyle, Spacer, KeepTogether, NextPageTemplate, CondPageBreak)
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_RIGHT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FONTS = os.path.join(ROOT, 'scripts', 'fonts')
OUT = os.path.join(ROOT, 'public', 'downloads')
os.makedirs(OUT, exist_ok=True)
items = json.load(open(os.path.join(ROOT, 'data', 'menu.json')))

GREEN, DEEP, BRASS, SAFFRON, IVORY, INK = (HexColor(x) for x in
    ('#0E3B2F', '#082A21', '#C9A24B', '#E8912D', '#FBF5E9', '#1F2A24'))
VEG, NONVEG = HexColor('#2E8B3E'), HexColor('#B3261E')

pdfmetrics.registerFont(TTFont('Rozha', os.path.join(FONTS, 'RozhaOne-Regular.ttf')))
pdfmetrics.registerFont(TTFont('Mukta', os.path.join(FONTS, 'Mukta-Regular.ttf')))
pdfmetrics.registerFont(TTFont('Mukta-SB', os.path.join(FONTS, 'Mukta-SemiBold.ttf')))
pdfmetrics.registerFont(TTFont('Mukta-B', os.path.join(FONTS, 'Mukta-Bold.ttf')))

SECTIONS = [
    ('biryani', 'Biryani & Pulao', 'Slow-cooked on dum. Small serves 1, large serves 1–2.'),
    ('tawa', 'Tawa Pulao & Anda Rice', 'Tossed to order on the tawa.'),
    ('starters', 'Kebabs & Starters', None),
    ('rolls', 'Rolls', None),
    ('combos', 'Combo Meals', 'Biryani, kebabs, gulab jamun and a drink in one order.'),
    ('bulk', 'Party & Bulk Orders', 'Priced per kg. Share your guest count and we will suggest quantities.'),
    ('beverages', 'Beverages', None),
]
PHONE, ADDRESS = '+91 96196 11561', 'Second Floor, A-Wing, Express Zone, Malad East, Mumbai 400097'
rs = lambda v: f'\u20b9{v:,}' if v else '–'

W, H = A4
M = 16 * mm

def header(c, doc):
    c.saveState()
    c.setFillColor(IVORY); c.rect(0, 0, W, H, stroke=0, fill=1)
    c.setFillColor(GREEN); c.rect(0, H - 62 * mm, W, 62 * mm, stroke=0, fill=1)
    logo = os.path.join(ROOT, 'public', 'logo.png')
    c.drawImage(logo, M, H - 52 * mm, 40 * mm, 40 * mm, mask='auto')
    c.setFillColor(IVORY); c.setFont('Rozha', 34)
    c.drawString(M + 48 * mm, H - 30 * mm, "Daddu's Biryani")
    c.setFillColor(BRASS); c.setFont('Mukta-SB', 13)
    c.drawString(M + 48 * mm, H - 38 * mm, 'Zayqo ki Kahani  |  Lucknowi, Hyderabadi, Kolkata & Mumbai dum biryani')
    c.setFillColor(IVORY); c.setFont('Mukta', 10)
    c.drawString(M + 48 * mm, H - 46 * mm, f'Order & bulk enquiries: {PHONE}  (call or WhatsApp)')
    c.setStrokeColor(BRASS); c.setLineWidth(1.2); c.line(0, H - 62 * mm, W, H - 62 * mm)
    footer(c, doc)
    c.restoreState()

def later(c, doc):
    c.saveState()
    c.setFillColor(IVORY); c.rect(0, 0, W, H, stroke=0, fill=1)
    c.setFillColor(GREEN); c.rect(0, H - 14 * mm, W, 14 * mm, stroke=0, fill=1)
    c.setFillColor(IVORY); c.setFont('Rozha', 13); c.drawString(M, H - 9.5 * mm, "Daddu's Biryani")
    c.setFillColor(BRASS); c.setFont('Mukta-SB', 10); c.drawRightString(W - M, H - 9.5 * mm, PHONE)
    footer(c, doc)
    c.restoreState()

def footer(c, doc):
    c.setStrokeColor(BRASS); c.setLineWidth(0.6); c.line(M, 14 * mm, W - M, 14 * mm)
    c.setFillColor(INK); c.setFont('Mukta', 8.5)
    c.drawString(M, 9 * mm, ADDRESS)
    c.drawRightString(W - M, 9 * mm, f'Page {doc.page}   |   Prices in INR, taxes as applicable')
    c.setFillColor(VEG); c.rect(M, 4 * mm, 2.6 * mm, 2.6 * mm, stroke=0, fill=1)
    c.setFillColor(INK); c.drawString(M + 4 * mm, 4.2 * mm, 'Veg')
    c.setFillColor(NONVEG); c.rect(M + 14 * mm, 4 * mm, 2.6 * mm, 2.6 * mm, stroke=0, fill=1)
    c.setFillColor(INK); c.drawString(M + 18 * mm, 4.2 * mm, 'Non-veg')

name_st = ParagraphStyle('n', fontName='Mukta-SB', fontSize=10.5, leading=13, textColor=INK)
note_st = ParagraphStyle('t', fontName='Mukta', fontSize=8.5, leading=10.5, textColor=HexColor('#5B6B61'))
price_st = ParagraphStyle('p', fontName='Mukta-B', fontSize=10.5, leading=13, textColor=GREEN, alignment=TA_RIGHT)
head_st = ParagraphStyle('h', fontName='Rozha', fontSize=20, leading=24, textColor=GREEN, spaceBefore=4)
sub_st = ParagraphStyle('s', fontName='Mukta', fontSize=9.5, leading=12, textColor=HexColor('#5B6B61'))
col_st = ParagraphStyle('c', fontName='Mukta-SB', fontSize=8.5, leading=10, textColor=BRASS, alignment=TA_RIGHT)

def dot(item):
    if 'veg' not in item: return ''
    color = '#2E8B3E' if item['veg'] else '#B3261E'
    return f'<font name="ZapfDingbats" size="7" color="{color}">n</font>&nbsp;&nbsp;'

story = [NextPageTemplate('later')]
content_w = W - 2 * M
for key, title, sub in SECTIONS:
    rows = [i for i in items if i['section'] == key]
    if not rows: continue
    has_sl = any('small' in i['prices'] or 'large' in i['prices'] for i in rows)
    has_kg = any('perKg' in i['prices'] for i in rows)
    cols = []
    if has_sl: cols += [('small', 'Small'), ('large', 'Large')]
    if has_kg: cols += [('perKg', 'Per kg')]
    if not cols or any('price' in i['prices'] for i in rows): cols = cols or [('price', 'Price')]
    pw = 20 * mm
    widths = [content_w - pw * len(cols)] + [pw] * len(cols)
    head = [KeepTogether([CondPageBreak(40 * mm), Paragraph(title, head_st)] + ([Paragraph(sub, sub_st)] if sub else []) + [Spacer(1, 3 * mm)])]
    data = [[''] + [Paragraph(lbl, col_st) for _, lbl in cols]]
    for it in rows:
        cell = [Paragraph(dot(it) + it['name'], name_st)]
        extra = it.get('note') or ''
        if it.get('includes'): extra = ', '.join(it['includes']) + (f"  (serves {it['serves']})" if it.get('serves') else '')
        if extra: cell.append(Paragraph(extra, note_st))
        pr = it['prices']
        vals = []
        for k, _ in cols:
            v = pr.get(k) if k != 'price' else pr.get('price')
            if k == 'small' and 'price' in pr and len(cols) > 1: v = pr['price']
            vals.append(Paragraph(rs(v), price_st))
        data.append([cell] + vals)
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LINEBELOW', (0, 1), (-1, -1), 0.4, HexColor('#E4D6B5')),
        ('LINEBELOW', (0, 0), (-1, 0), 0.9, BRASS),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 2),
    ]))
    story += head + [t, Spacer(1, 8 * mm)]

pdf_path = os.path.join(OUT, 'daddus-biryani-menu.pdf')
doc = BaseDocTemplate(pdf_path, pagesize=A4, title="Daddu's Biryani Menu", author="Daddu's Biryani")
doc.addPageTemplates([
    PageTemplate('first', [Frame(M, 18 * mm, W - 2 * M, H - 62 * mm - 24 * mm, id='f1')], onPage=header),
    PageTemplate('later', [Frame(M, 18 * mm, W - 2 * M, H - 14 * mm - 24 * mm, id='f2')], onPage=later),
])
doc.build(story)

# ---------- Excel ----------
wb = openpyxl.Workbook()
ws = wb.active; ws.title = 'Menu'
thin = Side(style='thin', color='E4D6B5')
ws['A1'] = "Daddu's Biryani – Menu"; ws['A1'].font = Font(name='Arial', size=18, bold=True, color='0E3B2F')
ws['A2'] = f'Call / WhatsApp {PHONE}  |  {ADDRESS}'; ws['A2'].font = Font(name='Arial', size=10, color='5B6B61')
ws['A3'] = 'Prices in INR. Taxes as applicable. Small serves 1, large serves 1–2.'; ws['A3'].font = Font(name='Arial', size=9, italic=True, color='5B6B61')
hdr = ['Section', 'Item', 'Veg / Non-veg', 'Small (₹)', 'Large (₹)', 'Per kg (₹)', 'Price (₹)', 'Details']
r = 5
for c, h in enumerate(hdr, 1):
    cell = ws.cell(r, c, h)
    cell.font = Font(name='Arial', bold=True, color='FBF5E9'); cell.fill = PatternFill('solid', fgColor='0E3B2F')
    cell.alignment = Alignment(vertical='center')
titles = {k: t for k, t, _ in SECTIONS}
for key, _, _ in SECTIONS:
    for it in [i for i in items if i['section'] == key]:
        r += 1
        pr = it['prices']
        details = it.get('note') or ''
        if it.get('includes'): details = ', '.join(it['includes']) + (f" (serves {it['serves']})" if it.get('serves') else '')
        vals = [titles[key], it['name'], '' if 'veg' not in it else ('Veg' if it['veg'] else 'Non-veg'),
                pr.get('small'), pr.get('large'), pr.get('perKg'), pr.get('price'), details]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(name='Arial', size=10, color='1F2A24'); cell.border = Border(bottom=thin)
            if 4 <= c <= 7: cell.number_format = '#,##0'; cell.alignment = Alignment(horizontal='right')
            if c == 8: cell.alignment = Alignment(wrap_text=True, vertical='top')
for col, w in zip('ABCDEFGH', [22, 46, 13, 11, 11, 11, 11, 60]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A6'
ws.auto_filter.ref = f'A5:H{r}'
xlsx_path = os.path.join(OUT, 'daddus-biryani-menu.xlsx')
wb.save(xlsx_path)
print('wrote', pdf_path, os.path.getsize(pdf_path)//1024, 'KB;', xlsx_path)
